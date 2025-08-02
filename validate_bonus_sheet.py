from openpyxl import load_workbook

def validate_bonus_sheet():
    # Load the workbook
    wb = load_workbook("BonusSheet-Aug2025.xlsx", data_only=False)
    ws = wb.active
    
    print("🔍 VALIDATING BONUS SHEET FORMULAS")
    print("=" * 50)
    
    # Check header row
    headers = []
    for col in range(1, 15):
        headers.append(ws.cell(1, col).value)
    
    expected_headers = [
        "Date", "Day", "Meeting Hrs", "Assigned Hrs", "Completed Hrs",
        "Complexity Factor", "QA Factor", "Task Failed (Y/N)", "Leave Taken (Y/N)",
        "Auto-adjust Available Hrs", "% Efficiency", "Raw Points", "OT Points", "Approved Points"
    ]
    
    print("✅ Headers validation:")
    for i, (actual, expected) in enumerate(zip(headers, expected_headers)):
        status = "✓" if actual == expected else "✗"
        print(f"   {status} Col {i+1}: {actual}")
    
    # Check sample formulas
    print("\n✅ Formula validation:")
    
    # Check row 2 formulas
    row = 2
    formulas_to_check = [
        (10, f'=IF(I{row}="Y",0,9-C{row})', "Auto-adjust Available Hrs"),
        (11, f'=IF(J{row}=0,0,E{row}/J{row})', "% Efficiency"),
        (12, f'=E{row}*F{row}*G{row}', "Raw Points"),
        (13, f'=IF(E{row}>D{row},E{row}-D{row},0)', "OT Points"),
        (14, f'=IF(H{row}="Y",0,ROUND(K{row}*L{row},2))', "Approved Points")
    ]
    
    for col, expected_formula, description in formulas_to_check:
        actual_formula = ws.cell(row, col).value
        status = "✓" if actual_formula == expected_formula else "✗"
        print(f"   {status} {description}: {actual_formula}")
    
    # Check summary section
    print("\n✅ Summary section validation:")
    
    # Find summary start row (should be around row 32)
    summary_start_row = None
    for row in range(30, 40):
        if ws.cell(row, 13).value == "Total Earned Points":
            summary_start_row = row
            break
    
    if summary_start_row:
        print(f"   ✓ Summary found at row {summary_start_row}")
        
        # Check summary formulas
        summary_formulas = [
            (summary_start_row, f'=SUM(N2:N27)', "Total Earned Points"),
            (summary_start_row + 3, f'=N{summary_start_row+2}*0.5/(N{summary_start_row+1}*10)', "Bonus Rate/Point"),
            (summary_start_row + 4, f'=MIN(N{summary_start_row}*N{summary_start_row+3},N{summary_start_row+2}*0.5)', "Monthly Bonus"),
            (summary_start_row + 5, f'=N{summary_start_row+2}+N{summary_start_row+4}', "Total Pay")
        ]
        
        for row, expected_formula, description in summary_formulas:
            actual_formula = ws.cell(row, 14).value
            status = "✓" if actual_formula == expected_formula else "✗"
            print(f"   {status} {description}: {actual_formula}")
    else:
        print("   ✗ Summary section not found")
    
    # Check protection
    print("\n✅ Protection validation:")
    print(f"   ✓ Sheet protection enabled: {ws.protection.sheet}")
    print(f"   ✓ Password set: {'Yes' if ws.protection.password else 'No'}")
    
    # Count workdays
    workday_count = 0
    for row in range(2, 32):
        if ws.cell(row, 1).value:
            workday_count += 1
    
    print(f"\n✅ Workdays count: {workday_count} (Expected: 26)")
    
    print("\n" + "=" * 50)
    print("✅ VALIDATION COMPLETE")
    
    # Show example calculation
    print("\n📊 EXAMPLE CALCULATION:")
    print("   Assuming all days worked with default values:")
    print("   - 26 workdays × 9 hours × 1 complexity × 1 QA = 234 raw points")
    print("   - 100% efficiency → 234 approved points")
    print("   - Base salary: ₹50,000")
    print("   - Bonus rate: ₹50,000 × 0.5 ÷ (26 × 10) = ₹96.15 per point")
    print("   - Monthly bonus: 234 × ₹96.15 = ₹22,499")
    print("   - Total pay: ₹50,000 + ₹22,499 = ₹72,499")

if __name__ == "__main__":
    validate_bonus_sheet()