from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment
from openpyxl.styles.protection import Protection
from openpyxl.utils import get_column_letter
from datetime import datetime, timedelta

# Initialize workbook and sheet
wb = Workbook()
ws = wb.active
ws.title = "TestEmployee"

# Headers and summary labels
headers = [
    "Date", "Day", "Meeting Hrs", "Assigned Hrs", "Completed Hrs",
    "Complexity Factor", "QA Factor", "Task Failed (Y/N)", "Leave Taken (Y/N)",
    "Auto-adjust Available Hrs", "% Efficiency", "Raw Points", "OT Points", "Approved Points"
]
summary_labels = [
    "Total Earned Points", "Workdays (W)", "Base Salary", "Bonus Rate/Point",
    "Monthly Bonus", "Total Pay (Basic + Bonus)"
]

# Styling
header_font = Font(bold=True)
header_fill = PatternFill(start_color="FFD966", end_color="FFD966", fill_type="solid")
summary_fill = PatternFill(start_color="F4CCCC", end_color="F4CCCC", fill_type="solid")
alignment_center = Alignment(horizontal="center", vertical="center")

# Add headers
for col_num, header in enumerate(headers, 1):
    cell = ws.cell(row=1, column=col_num, value=header)
    cell.font = header_font
    cell.fill = header_fill
    cell.alignment = alignment_center

# Generate example data for workdays in August 2025 (Mon-Sat)
start_date = datetime(2025, 8, 1)
workdays = []
for i in range(31):
    day = start_date + timedelta(days=i)
    if day.weekday() < 6:  # Mon-Sat (0=Monday, 5=Saturday)
        workdays.append(day)

# Fill data rows
for row_num, date in enumerate(workdays, 2):
    day_name = date.strftime("%a")
    ws.cell(row=row_num, column=1, value=date.strftime("%Y-%m-%d"))
    ws.cell(row=row_num, column=2, value=day_name)
    ws.cell(row=row_num, column=3, value=0)  # Meeting Hrs
    ws.cell(row=row_num, column=4, value=9)  # Assigned Hrs
    ws.cell(row=row_num, column=5, value=9)  # Completed Hrs
    ws.cell(row=row_num, column=6, value=1)  # Complexity Factor
    ws.cell(row=row_num, column=7, value=1)  # QA Factor
    ws.cell(row=row_num, column=8, value="N")  # Task Failed
    ws.cell(row=row_num, column=9, value="N")  # Leave Taken
    
    # Auto-adjust Available Hrs = 9 - Meeting Hrs if not on leave
    ws.cell(row=row_num, column=10, value=f'=IF(I{row_num}="Y",0,9-C{row_num})')
    # % Efficiency = Completed / Auto-adjust Available
    ws.cell(row=row_num, column=11, value=f'=IF(J{row_num}=0,0,E{row_num}/J{row_num})')
    # Raw Points = Completed * Complexity * QA
    ws.cell(row=row_num, column=12, value=f'=E{row_num}*F{row_num}*G{row_num}')
    # OT Points = Completed - Assigned if > Assigned
    ws.cell(row=row_num, column=13, value=f'=IF(E{row_num}>D{row_num},E{row_num}-D{row_num},0)')
    # Approved Points = 0 if Task Failed, else Raw Points + OT
    ws.cell(row=row_num, column=14, value=f'=IF(H{row_num}="Y",0,ROUND(L{row_num}+M{row_num},2))')

# Summary calculations
summary_start_row = len(workdays) + 4
workdays_count = len(workdays)
base_salary = 50000

# Fixed summary formulas
summary_values = [
    f'=SUM(N2:N{len(workdays)+1})',        # Total Earned Points
    workdays_count,                        # Workdays (W)
    base_salary,                           # Base Salary
    f'=N{summary_start_row+2}*0.5/N{summary_start_row+1}',  # Bonus Rate per point (Base Salary * 0.5 / Workdays)
    f'=N{summary_start_row}*N{summary_start_row+3}',        # Monthly Bonus (Total Points * Rate)
    f'=N{summary_start_row+2}+N{summary_start_row+4}'       # Total Pay (Base + Bonus)
]

# Add summary section
for i, (label, val) in enumerate(zip(summary_labels, summary_values)):
    label_cell = ws.cell(row=summary_start_row + i, column=13, value=label)
    label_cell.font = Font(bold=True)
    label_cell.fill = summary_fill
    
    value_cell = ws.cell(row=summary_start_row + i, column=14, value=val)

# Set column widths
for col in range(1, len(headers) + 1):
    ws.column_dimensions[get_column_letter(col)].width = 18

# Protect sheet and unlock specific editable cells
ws.protection.sheet = True
ws.protection.password = "secure123"

# Unlock input cells (editable ones)
for row in range(2, len(workdays) + 2):
    for col in [3, 4, 5, 6, 7, 8, 9]:  # Meeting Hrs, Assigned Hrs, Completed Hrs, Complexity, QA, Task Failed, Leave Taken
        ws.cell(row=row, column=col).protection = Protection(locked=False)

# Unlock Base Salary input cell
ws.cell(row=summary_start_row + 2, column=14).protection = Protection(locked=False)

# Save the file to current directory
file_path = "EmployeeBonusSheet-Aug2025.xlsx"
wb.save(file_path)

print(f"Excel file created successfully: {file_path}")
print(f"Total workdays in August 2025: {len(workdays)}")
print("Editable fields:")
print("- Meeting Hrs, Assigned Hrs, Completed Hrs")
print("- Complexity Factor, QA Factor")
print("- Task Failed (Y/N), Leave Taken (Y/N)")
print("- Base Salary in summary section")